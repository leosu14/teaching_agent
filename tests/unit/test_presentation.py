"""Presentation slice units: slide plan schema and validator, builder, mock and PPTX renderers, config and theme.

The deck here is a biology lesson on purpose: nothing in the presentation layer assumes a language course.
"""

from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime, timezone

import pytest
from pptx import Presentation as PptxDocument
from pptx.util import Emu
from pydantic import ValidationError

from app.providers.presentation.base import PresentationRenderError
from app.providers.presentation.mock import MockPresentationRenderer
from app.providers.presentation.pptx_renderer import EMU_PER_PT, PptxPresentationRenderer, layout_regions
from app.schemas.artifact import StoredObject
from app.schemas.presentation import (
    ImageBlock,
    PresentationConfig,
    PresentationTheme,
    SlideDeckPlan,
    SlideLayout,
    SlidePlan,
    SlidePlanValidationRequest,
    TableBlock,
    ThemeFonts,
)
from app.schemas.research import (
    Citation,
    Evidence,
    ExtractionMetadata,
    ResearchBundle,
    ResearchObjective,
    ResearchTarget,
    Source,
)
from app.schemas.visual import GenerationAttribution, ImageAsset, ImageValidationReport, SearchAttribution
from app.schemas.visual import ImageLicense, SelectionRecord, SelectionSignals
from app.tools.presentation.builder import PresentationBuilder, PresentationBuildError
from app.tools.presentation.validation import SlidePlanValidator
from app.utils.images import encode_png

NOW = datetime(2026, 10, 1, tzinfo=timezone.utc)
PICTURE = 13


def research() -> ResearchBundle:
    sources = [Source(source_id=f"src{i}", url=f"https://example.org/{i}", canonical_url=f"https://example.org/{i}",
                      title=f"Plant biology {i}", publisher="Open Textbook", retrieved_at=NOW, retrieved_via="web",
                      provider="mock") for i in (1, 2)]
    evidence = [Evidence(evidence_id=f"ev{i}", source_id=f"src{i}", target_id="photosynthesis",
                         text="Plants turn light into chemical energy.", relevance=0.9,
                         extraction=ExtractionMetadata(method="quote", extractor="mock", extracted_at=NOW))
                for i in (1, 2)]
    citations = [Citation(citation_id=f"c{i}", evidence_id=f"ev{i}", source_id=f"src{i}", title=f"Plant biology {i}",
                          url=f"https://example.org/{i}", publisher="Open Textbook", retrieved_at=NOW) for i in (1, 2)]
    return ResearchBundle(
        research_id="res1", status="complete", sources=sources, evidence=evidence, citations=citations,
        generated_at=NOW, objective=ResearchObjective(
            description="d", subject="biology", topic="photosynthesis",
            targets=[ResearchTarget(target_id="photosynthesis", name="Photosynthesis")]),
    )


def image(name: str, origin: str, width: int = 1600, height: int = 900) -> tuple[ImageAsset, bytes]:
    data = encode_png(width, height, [(10, 120, 40), (200, 200, 30), (20, 40, 160)] if origin == "search"
                      else [(90, 20, 90), (240, 240, 240), (0, 0, 0)])
    checksum = hashlib.sha256(data).hexdigest()
    obj = StoredObject(uri=f"file:///store/{checksum}.png", checksum=checksum, media_type="image/png",
                       size_bytes=len(data))
    report = ImageValidationReport(valid=True, validator="test")
    if origin == "search":
        attribution = SearchAttribution(provider="mock", image_id=name, provider_image_id=name,
                                        image_url="https://img.example.org/leaf.png",
                                        source_url="https://img.example.org/leaf", title="A leaf in sunlight",
                                        creator="A. Botanist", license=ImageLicense(name="CC BY 4.0"),
                                        retrieved_at=NOW)
        signals = SelectionSignals(relevance=1, aspect=1, visual_type=1, license=1, source_quality=1, resolution=1)
        selection = SelectionRecord(selector="t", rank=1, score=1.0, signals=signals, candidates=1)
    else:
        attribution = GenerationAttribution(provider="mock-gen", model="m1", generated_at=NOW, prompt="p",
                                            prompt_hash="h", width=width, height=height)
        selection = None
    asset = ImageAsset(asset_id=f"img_{name}", visual_id=name, plan_id="vp1", lesson_section_id="sec_light",
                       visual_type="photo" if origin == "search" else "diagram", purpose="show it",
                       description=f"{name} alt text", origin=origin, object=obj, width=width, height=height,
                       format="image/png", attribution=attribution, selection=selection, validation=report)
    return asset, data


def slide(order: int, slide_type: str, layout: str, title: str, blocks=(), **kw) -> dict:
    return {"slide_id": f"s{order}", "order": order, "slide_type": slide_type, "layout": layout, "title": title,
            "content_blocks": list(blocks), **kw}


def deck_dict() -> dict:
    return {
        "deck_id": "deck1", "title": "Photosynthesis", "language": "en", "level": "intro", "topic": "photosynthesis",
        "objective": "Explain how plants make sugar from light.", "slides": [
            slide(1, "title", "title", "Photosynthesis", subtitle="Biology | Introductory"),
            slide(2, "objectives", "title_content", "Objectives",
                  [{"kind": "bullets", "items": ["Name the inputs", "Name the outputs"]}]),
            slide(3, "explanation", "image_text", "Light reactions",
                  [{"kind": "bullets", "items": ["Light is absorbed by chlorophyll"]},
                   {"kind": "image", "artifact_id": "art_photo", "caption": "A leaf in sunlight"}],
                  visual_refs=["art_photo"], citation_refs=["c1"], section_refs=["sec_light"],
                  speaker_notes="Point at the leaf."),
            slide(4, "comparison", "two_column", "Inputs and outputs",
                  [{"kind": "table", "headers": ["Inputs", "Outputs"], "rows": [["CO2", "O2"], ["Water", "Glucose"]]},
                   {"kind": "text", "text": "Mass is conserved."}], citation_refs=["c2"]),
            slide(5, "vocabulary", "title_content", "Key terms",
                  [{"kind": "vocabulary", "entries": [{"term": "Chlorophyll", "meaning": "Green pigment"}]}]),
            slide(6, "explanation", "full_image", "The whole cycle",
                  [{"kind": "image", "artifact_id": "art_diagram"}, {"kind": "text", "text": "The Calvin cycle"}],
                  visual_refs=["art_diagram"], section_refs=["sec_light"]),
            slide(7, "exercise", "exercise", "Check",
                  [{"kind": "question", "question_id": "q1", "prompt": "What gas is released?",
                    "choices": ["Oxygen", "Nitrogen"]}]),
            slide(8, "answer", "exercise", "Answer", [{"kind": "answer", "question_id": "q1", "answer": "Oxygen"}]),
            slide(9, "summary", "summary", "Summary", [{"kind": "text", "text": "Light becomes sugar."}]),
            slide(10, "references", "title_content", "References", [{"kind": "citations", "citation_ids": ["c1", "c2"]}]),
        ],
    }


def request(deck: dict | None = None, **overrides) -> SlidePlanValidationRequest:
    return SlidePlanValidationRequest(**{
        "deck": deck or deck_dict(), "section_ids": ["sec_light"], "citation_ids": ["c1", "c2"],
        "image_artifact_ids": ["art_photo", "art_diagram"], "question_ids": ["q1"], **overrides})


def codes(deck: dict, **overrides) -> set[str]:
    report = SlidePlanValidator().validate(request(deck, **overrides))
    return {e.code for e in report.errors}


def images() -> tuple[dict[str, ImageAsset], dict[str, bytes]]:
    photo, photo_bytes = image("photo", "search")
    diagram, diagram_bytes = image("diagram", "generated", 1280, 720)
    return ({"art_photo": photo, "art_diagram": diagram},
            {photo.object.uri: photo_bytes, diagram.object.uri: diagram_bytes})


def built(config: PresentationConfig | None = None):
    assets, media = images()
    deck = SlideDeckPlan.model_validate(deck_dict())
    return PresentationBuilder().build(deck, config=config or PresentationConfig(), research=research(),
                                       images=assets), media


# --- schema ----------------------------------------------------------------------------------------


def test_slides_are_structured_blocks_not_raw_strings_or_html() -> None:
    with pytest.raises(ValidationError):
        SlidePlan.model_validate({**deck_dict()["slides"][1], "content_blocks": ["<ul><li>raw html</li></ul>"]})
    with pytest.raises(ValidationError):
        SlidePlan.model_validate({**deck_dict()["slides"][1], "slide_type": "poster"})
    with pytest.raises(ValidationError):  # an image block holds only the artifact id, never image metadata
        ImageBlock.model_validate({"artifact_id": "art_photo", "url": "https://x/y.png", "width": 10})
    with pytest.raises(ValidationError):
        TableBlock(headers=["a", "b"], rows=[["only one"]])
    with pytest.raises(ValidationError):
        SlidePlan.model_validate({**deck_dict()["slides"][1], "title": "   "})


def test_config_defaults_to_16_9_and_rejects_mismatched_sizes() -> None:
    config = PresentationConfig()
    assert (config.aspect_ratio, config.width, config.height) == ("16:9", 960, 540)
    assert PresentationConfig.for_aspect("4:3").width == 720
    with pytest.raises(ValidationError):
        PresentationConfig(aspect_ratio="16:9", width=720, height=540)


# --- validator ------------------------------------------------------------------------------------


def test_a_well_formed_deck_validates() -> None:
    report = SlidePlanValidator().validate(request())
    assert report.valid and report.errors == [] and report.deck is not None and report.deck_id == "deck1"
    assert {"unique_slide_ids", "ordering", "image_refs", "citation_refs", "section_refs"} <= set(report.checks)


@pytest.mark.parametrize("mutate, expected", [
    (lambda d: d["slides"][2].update(slide_id="s2"), "duplicate_slide_id"),
    (lambda d: d["slides"][2].update(order=7), "slide_order"),
    (lambda d: d["slides"][2].update(section_refs=["sec_invented"]), "unknown_section"),
    (lambda d: d["slides"][2].update(citation_refs=["c99"]), "unknown_citation"),
    (lambda d: d["slides"][9]["content_blocks"][0].update(citation_ids=["c1", "c42"]), "unknown_citation"),
    (lambda d: (d["slides"][2].update(visual_refs=["art_gone"]),
                d["slides"][2]["content_blocks"][1].update(artifact_id="art_gone")), "unknown_image"),
    (lambda d: d["slides"][2].update(visual_refs=[]), "image_not_declared"),
    (lambda d: d["slides"][1].update(visual_refs=["art_photo"]), "unplaced_visual"),
    (lambda d: d["slides"][6]["content_blocks"][0].update(question_id="q9"), "unknown_question"),
    (lambda d: d["slides"].__setitem__(slice(6, 8), [{**d["slides"][7], "order": 7, "slide_id": "s7"},
                                                     {**d["slides"][6], "order": 8, "slide_id": "s8"}]),
     "answer_before_question"),
    (lambda d: d["slides"][2].update(layout="title_content"), "layout_mismatch"),
    (lambda d: d["slides"][4].update(layout="image_text"), "layout_mismatch"),
    (lambda d: d["slides"][4].update(slide_type="exercise"), "type_mismatch"),
    (lambda d: d["slides"][9].update(content_blocks=[{"kind": "text", "text": "see the web"}]), "type_mismatch"),
    (lambda d: d["slides"][1].update(content_blocks=[{"kind": "text", "text": "word " * 19}] * 7), "overcrowded"),
    (lambda d: d["slides"][1].update(content_blocks=[]), "empty_slide"),
    (lambda d: d["slides"][0].update(slide_type="summary"), "missing_title_slide"),
    (lambda d: d["slides"][1].update(slide_type="poster"), "invalid_schema"),
    (lambda d: d["slides"][1].update(content_blocks=["<p>raw</p>"]), "invalid_schema"),
])
def test_malformed_decks_fail_with_structured_errors(mutate, expected) -> None:
    deck = deck_dict()
    mutate(deck)
    assert expected in codes(deck)


def test_too_many_slides_and_issue_descriptions() -> None:
    report = SlidePlanValidator().validate(request(max_slides=5))
    [issue] = report.errors
    assert issue.code == "too_many_slides" and not report.valid
    assert issue.describe().startswith("[too_many_slides] deck:")


# --- builder ----------------------------------------------------------------------------------------


def test_builder_resolves_citations_and_images_into_layout_regions() -> None:
    presentation, _ = built()
    assert [s.order for s in presentation.slides] == list(range(1, 11))
    assert [s.slide_id for s in presentation.slides] == [f"s{i}" for i in range(1, 11)]
    # Citations are numbered by first appearance and resolved Citation -> Evidence -> Source.
    assert [(r.number, r.citation_id, r.evidence_id, r.source_id) for r in presentation.references] == [
        (1, "c1", "ev1", "src1"), (2, "c2", "ev2", "src2")]
    explanation = presentation.slides[2]
    image_el = explanation.element("image")[0]
    assert (image_el.region, image_el.artifact_id, image_el.alt_text) == ("image", "art_photo", "photo alt text")
    assert image_el.checksum == images()[0]["art_photo"].object.checksum
    footer = explanation.element("footer")[0]
    assert footer.citation_ids == ["c1"] and footer.image_artifact_ids == ["art_photo"]
    assert "[1] Plant biology 1" in footer.text and "A leaf in sunlight" in footer.text
    assert {e.region for e in presentation.slides[3].elements if e.kind != "footer"} >= {"left", "right"}
    assert presentation.slides[5].element("image")[0].region == "image"
    assert presentation.slides[5].element("text")[0].region == "caption"
    assert "Generated image (mock-gen/m1)" in presentation.slides[5].element("footer")[0].text
    assert presentation.slides[4].element("table")[0].headers == ["Term", "Meaning"]
    assert presentation.slides[6].element("bullets")[0].numbered
    refs = presentation.slides[9].element("bullets")[0].items
    assert refs[0].startswith("[1] Plant biology 1. Open Textbook. https://example.org/1")
    assert presentation.image_artifact_ids() == ["art_photo", "art_diagram"]
    again, _ = built()
    assert again == presentation  # deterministic, including the presentation id


def test_builder_refuses_unresolvable_references() -> None:
    assets, _ = images()
    deck = SlideDeckPlan.model_validate(deck_dict())
    with pytest.raises(PresentationBuildError, match="art_diagram"):
        PresentationBuilder().build(deck, config=PresentationConfig(), research=research(),
                                    images={"art_photo": assets["art_photo"]})
    bad = deck_dict()
    bad["slides"][2]["citation_refs"] = ["c77"]
    with pytest.raises(PresentationBuildError, match="c77"):
        PresentationBuilder().build(SlideDeckPlan.model_validate(bad), config=PresentationConfig(),
                                    research=research(), images=assets)


# --- renderers -------------------------------------------------------------------------------------


async def test_mock_renderer_is_deterministic_and_complete() -> None:
    presentation, _ = built()
    first = await MockPresentationRenderer().render(presentation)
    second = await MockPresentationRenderer().render(presentation)
    assert first.content == second.content and first.media_type == "application/json"
    assert (first.slides, first.elements) == (10, presentation.element_count())
    assert first.image_artifact_ids == ["art_photo", "art_diagram"]
    doc = json.loads(first.content)
    assert [s["order"] for s in doc["slides"]] == list(range(1, 11))
    assert [e["artifact_id"] for s in doc["slides"] for e in s["elements"] if e["kind"] == "image"] == \
        ["art_photo", "art_diagram"]
    assert [r["citation_id"] for r in doc["references"]] == ["c1", "c2"]
    assert doc["slides"][2]["citation_ids"] == ["c1"]


async def test_pptx_renderer_writes_a_real_reproducible_file() -> None:
    presentation, media = built()
    rendered = await PptxPresentationRenderer(media=media.__getitem__).render(presentation)
    again = await PptxPresentationRenderer(media=media.__getitem__).render(presentation)
    assert rendered.content == again.content  # byte-reproducible: same presentation, same checksum
    assert rendered.media_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    doc = PptxDocument(io.BytesIO(rendered.content))
    assert len(doc.slides) == 10
    assert (doc.slide_width, doc.slide_height) == (Emu(960 * EMU_PER_PT), Emu(540 * EMU_PER_PT))
    assert [s.shapes.title.text for s in doc.slides] == [s.title for s in SlideDeckPlan.model_validate(deck_dict()).slides]
    assert doc.core_properties.title == "Photosynthesis"

    pictures = {i: [s for s in sl.shapes if s.shape_type == PICTURE] for i, sl in enumerate(doc.slides, start=1)}
    assert {i for i, p in pictures.items() if p} == {3, 6}
    assert hashlib.sha256(pictures[3][0].image.blob).hexdigest() == presentation.slides[2].element("image")[0].checksum
    assert pictures[3][0]._element.nvPicPr.cNvPr.get("descr") == "photo alt text"

    texts = {i: " ".join(s.text_frame.text for s in sl.shapes if s.has_text_frame) for i, sl in
             enumerate(doc.slides, start=1)}
    assert "Light is absorbed by chlorophyll" in texts[3] and "1. Oxygen" in texts[7]
    footer = next(s for s in doc.slides[2].shapes if s.name == "footer").text_frame.text
    assert "[1] Plant biology 1" in footer and "A leaf in sunlight" in footer
    assert next(s for s in doc.slides[2].shapes if s.name == "slide_number").text_frame.text == "3/10"
    assert doc.slides[2].notes_slide.notes_text_frame.text == "Point at the leaf."
    table = next(s for s in doc.slides[3].shapes if s.has_table).table
    assert [c.text for c in table.rows[0].cells] == ["Inputs", "Outputs"]

    # Every shape stays on the slide, and images never overlap the text next to them.
    for sl in doc.slides:
        for shape in sl.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= doc.slide_width and shape.top + shape.height <= doc.slide_height
    for sl in (doc.slides[2], doc.slides[3]):
        pic = next((s for s in sl.shapes if s.shape_type == PICTURE or s.has_table), None)
        body = [s for s in sl.shapes if s.has_text_frame and s.name.split(".")[-1].startswith(("bullets", "text"))]
        for shape in body:
            assert shape.left + shape.width <= pic.left or pic.left + pic.width <= shape.left or \
                shape.top + shape.height <= pic.top or pic.top + pic.height <= shape.top


async def test_pptx_renderer_never_places_bytes_that_do_not_match_the_asset() -> None:
    presentation, media = built()
    tampered = {uri: data + b"x" for uri, data in media.items()}
    with pytest.raises(PresentationRenderError, match="checksum"):
        await PptxPresentationRenderer(media=tampered.__getitem__).render(presentation)

    def missing(uri: str) -> bytes:
        raise OSError("gone")
    with pytest.raises(PresentationRenderError, match="cannot be read"):
        await PptxPresentationRenderer(media=missing).render(presentation)


async def test_theme_and_config_drive_the_geometry_and_fonts() -> None:
    config = PresentationConfig.for_aspect("4:3", theme=PresentationTheme(fonts=ThemeFonts(heading="Georgia",
                                                                                           body="Verdana")))
    presentation, media = built(config)
    rendered = await PptxPresentationRenderer(media=media.__getitem__).render(presentation)
    doc = PptxDocument(io.BytesIO(rendered.content))
    assert (doc.slide_width, doc.slide_height) == (Emu(720 * EMU_PER_PT), Emu(540 * EMU_PER_PT))
    title_run = doc.slides[1].shapes.title.text_frame.paragraphs[0].runs[0]
    assert title_run.font.name == "Georgia"
    body = next(s for s in doc.slides[1].shapes if ".bullets" in s.name)
    assert body.text_frame.paragraphs[0].runs[0].font.name == "Verdana"
    wide, narrow = layout_regions(SlideLayout.IMAGE_TEXT, PresentationConfig()), layout_regions(SlideLayout.IMAGE_TEXT, config)
    assert narrow["image"].w < wide["image"].w and narrow["body"].x == wide["body"].x == config.theme.spacing.margin
