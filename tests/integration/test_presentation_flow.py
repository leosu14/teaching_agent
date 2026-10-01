"""The presentation stages inside the lesson workflow: slide planning, validation gate, build, PPTX render, the
PRESENTATION artifact, visual placement, citations, review gating, failure policy, events, cost, resume and
idempotency."""

from __future__ import annotations

import hashlib
import io
import json

import pytest
from pptx import Presentation as PptxDocument

from app.config.settings import Settings
from app.providers.llm.mock import MockLLMProvider
from app.providers.llm.mock_responders import default_responders
from app.providers.presentation.base import PresentationRenderer, PresentationRenderError
from app.schemas.artifact import ArtifactType
from app.schemas.lesson import LessonContent
from app.schemas.presentation import (
    Presentation,
    PresentationArtifactMetadata,
    SlideDeckPlan,
    SlidePlanningInput,
)
from app.schemas.research import ResearchBundle
from app.schemas.task import TaskStatus
from app.schemas.visual import ImageAsset
from app.services.container import build_container
from tests.conftest import add_demo_learner, answers_for, run_lesson
from tests.integration.test_resume import SimulatedCrash, crash_after
from tests.integration.test_review_loop import always_reject
from tests.integration.test_visual_flow import GenerationFailing, SearchWithout

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PICTURE = 13
PRESENTATION_EVENTS = ["slide_planning.started", "slide_plan.created", "slide_plan.validated",
                       "presentation.build_started", "presentation.build_completed", "presentation.render_started",
                       "presentation.render_completed", "presentation.artifact_created"]
PROVIDERS = ("image_search_provider", "image_generation_provider", "presentation_renderer")


class FailingRenderer(PresentationRenderer):
    name, format, media_type = "broken", "pptx", PPTX

    async def render(self, presentation):
        raise PresentationRenderError("disk full while writing the package")


@pytest.fixture
def make_container(tmp_path):
    made = []

    def make(*, llm: MockLLMProvider | None = None, observers=(), data_dir=None, **kwargs):
        # A renderer instance is injected; the string form is the TA_PRESENTATION_RENDERER setting.
        providers = {k: kwargs.pop(k) for k in PROVIDERS if k in kwargs and not isinstance(kwargs[k], str)}
        c = build_container(Settings(data_dir=data_dir or tmp_path / f"d{len(made)}", log_json=False, **kwargs),
                            llm_providers={"mock": llm or MockLLMProvider(default_responders())},
                            observers=observers, **providers)
        made.append(c)
        return c

    yield make
    for c in made:
        c.close()


def arts_of(container, task) -> dict:
    return {a.name: a for a in container.task_service.artifacts(task.task_id)}


def deck_of(container, task) -> SlideDeckPlan:
    return SlideDeckPlan.model_validate_json(container.artifacts.read(arts_of(container, task)["slide_plan"].artifact_id))


def pptx_of(container, task):
    return PptxDocument(io.BytesIO(container.artifacts.read(arts_of(container, task)["presentation"].artifact_id)))


def presentation_events(container, task) -> list:
    return [e for e in container.task_service.events(task.task_id)
            if e.type.startswith(("slide_planning.", "slide_plan.", "presentation."))]


async def test_approved_lesson_becomes_a_real_pptx_with_its_images_and_citations(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    container = make_container(llm=llm)
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    order = [n for n in task.workflow.execution_order if n in {
        "teach_review", "visual", "store_artifacts", "slide_plan", "validate_slide_plan", "store_slide_plan",
        "build_presentation", "render_presentation", "update_learner"}]
    assert order == ["teach_review", "visual", "store_artifacts", "slide_plan", "validate_slide_plan",
                     "store_slide_plan", "build_presentation", "render_presentation", "update_learner"]

    # The planner saw the approved lesson, its plan, the research citations and the image assets: ids, no bytes.
    planning = SlidePlanningInput.model_validate(next(r for r in llm.requests if r.agent_id == "slide_planner").input_payload)
    arts = arts_of(container, task)
    images = {a.artifact_id: a for a in arts.values() if a.type == ArtifactType.IMAGE_ASSET}
    research = ResearchBundle.model_validate_json(container.artifacts.read(arts["research_bundle"].artifact_id))
    assert {v.artifact_id for v in planning.visuals} == set(images)
    assert {c.citation_id for c in planning.citations} == {c.citation_id for c in research.citations}
    assert "file://" not in json.dumps(planning.model_dump(mode="json")["visuals"])
    assert planning.lesson.sections and planning.plan.objectives

    # Artifact graph: Lesson -> SLIDE_PLAN -> PRESENTATION, which also hangs off the lesson and its placed images.
    deck = deck_of(container, task)
    lesson_art, plan_art, pres_art = arts["lesson"], arts["slide_plan"], arts["presentation"]
    assert plan_art.type == ArtifactType.SLIDE_PLAN and plan_art.parent_ids == [lesson_art.artifact_id]
    assert pres_art.type == ArtifactType.PRESENTATION and pres_art.media_type == PPTX
    assert pres_art.parent_ids == [plan_art.artifact_id, lesson_art.artifact_id, *deck.image_artifact_ids()]
    assert {a.name for a in container.artifacts.lineage(pres_art.artifact_id)} >= {
        "slide_plan", "lesson", "lesson_plan", "research_bundle", "visual_plan", *(a.name for a in images.values())}
    data = container.artifacts.read(pres_art.artifact_id)
    assert hashlib.sha256(data).hexdigest() == pres_art.content_hash
    meta = PresentationArtifactMetadata.model_validate(pres_art.metadata)
    assert meta.slide_plan_artifact_id == plan_art.artifact_id and meta.deck_id == deck.deck_id
    assert meta.checksum == pres_art.content_hash and meta.media_type == PPTX and meta.renderer == "python-pptx"
    assert meta.object_key == f"objects/sha256/{meta.checksum[:2]}/{meta.checksum}.pptx"  # not a filesystem path
    assert meta.slides == len(deck.slides) and set(meta.image_artifact_ids) == set(images)
    assert {pres_art.artifact_id, plan_art.artifact_id} <= {a.artifact_id for a in task.result.artifacts}
    assert pres_art.artifact_id in task.artifact_ids

    # The PPTX itself, opened with python-pptx.
    doc = pptx_of(container, task)
    assert len(doc.slides) == len(deck.slides)
    assert [s.shapes.title.text for s in doc.slides] == [s.title for s in deck.slides]  # same order
    checksums = {ImageAsset.model_validate(a.metadata).object.checksum: aid for aid, a in images.items()}
    lesson = LessonContent.model_validate_json(container.artifacts.read(lesson_art.artifact_id))
    for planned, page in zip(deck.slides, doc.slides):
        placed = [checksums[hashlib.sha256(s.image.blob).hexdigest()] for s in page.shapes if s.shape_type == PICTURE]
        assert placed == [b.artifact_id for b in planned.blocks("image")]  # the stored IMAGE_ASSET bytes, as planned
        text = " ".join(s.text_frame.text for s in page.shapes if s.has_text_frame)
        for block in planned.blocks("bullets"):
            assert all(item in text for item in block.items)
        footer = next((s.text_frame.text for s in page.shapes if s.name == "footer"), "")
        for cid in planned.citation_refs:
            assert research.resolve(cid)[0].title in footer
        if planned.speaker_notes:
            assert page.notes_slide.notes_text_frame.text == planned.speaker_notes
    placed_all = [b.artifact_id for s in deck.slides for b in s.blocks("image")]
    assert sorted(placed_all) == sorted(images)  # closes PR #4's gap: every image asset is on a slide
    # Every section is presented and every slide citation is one the lesson's research produced.
    assert {r for s in deck.slides for r in s.section_refs} == {s.section_id for s in lesson.sections}
    assert set(deck.citation_ids()) == {c for s in lesson.sections for c in s.citations}


async def test_presentation_events_and_cost(make_container) -> None:
    container = make_container()
    task = await run_lesson(container)
    events = presentation_events(container, task)
    assert [e.type for e in events] == PRESENTATION_EVENTS
    created = events[-1].data
    assert created["reused"] is False and created["artifact_id"] == arts_of(container, task)["presentation"].artifact_id
    cost = task.cost
    assert cost.by_agent["slide_planner"].calls == 1 and cost.by_agent["slide_planner"].cost_usd > 0
    assert cost.by_agent["slide_planner"].usage.total_tokens > 0
    render = cost.by_service["presentation_render:python-pptx"]
    assert render.cost_usd is None and render.units["slides"] == len(deck_of(container, task).slides)
    assert "presentation_render" not in "".join(cost.by_agent) and "python-pptx" not in cost.by_model


async def test_planner_corrects_a_plan_that_fails_validation(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    container = make_container(llm=llm)
    # First proposal invents a citation id; the validator (through the ToolManager) sends it back.
    invented = {"title": "T", "slides": [
        {"slide_id": "s1", "order": 1, "slide_type": "title", "layout": "title", "title": "T"},
        {"slide_id": "s2", "order": 2, "slide_type": "summary", "layout": "summary", "title": "S",
         "content_blocks": [{"kind": "text", "text": "x"}], "citation_refs": ["c_invented"]}]}
    llm.inject("slide_planner", json.dumps(invented))
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    assert llm.calls["slide_planner"] == 2
    second = SlidePlanningInput.model_validate([r for r in llm.requests if r.agent_id == "slide_planner"][1].input_payload)
    assert any("unknown_citation" in c and "c_invented" in c for c in second.corrections)
    failed = [e for e in container.task_service.events(task.task_id)
              if e.type == "agent.validation_failed" and e.agent_id == "slide_planner"]
    assert len(failed) == 1 and "c_invented" in failed[0].data["error"]
    assert "c_invented" not in deck_of(container, task).citation_ids()


async def test_a_malformed_slide_plan_fails_the_task_before_anything_is_built(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    container = make_container(llm=llm)
    broken = {"title": "T", "slides": [
        {"slide_id": "s1", "order": 1, "slide_type": "title", "layout": "title", "title": "T"},
        {"slide_id": "s1", "order": 3, "slide_type": "explanation", "layout": "image_text", "title": "E",
         "content_blocks": [{"kind": "image", "artifact_id": "art_not_an_asset"}], "visual_refs": ["art_not_an_asset"]}]}
    llm.inject("slide_planner", *[json.dumps(broken)] * 3)  # every correction attempt stays broken
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED
    assert task.errors[-1].node_id == "validate_slide_plan"
    for code in ("duplicate_slide_id", "slide_order", "unknown_image"):
        assert code in task.errors[-1].message
    assert task.workflow.node_states["build_presentation"].status.value == "PENDING"
    names = set(arts_of(container, task))
    assert "presentation" not in names and "slide_plan" not in names and "lesson" in names
    failed = [e for e in presentation_events(container, task) if e.type == "presentation.failed"]
    assert failed and failed[-1].data["stage"] == "validation"
    assert {e["code"] for e in failed[-1].data["errors"]} >= {"duplicate_slide_id", "unknown_image"}


async def test_schema_invalid_slides_are_rejected_and_regenerated(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    container = make_container(llm=llm)
    html = {"title": "T", "slides": [{"slide_id": "s1", "order": 1, "slide_type": "title", "layout": "title",
                                      "title": "T", "content_blocks": ["<h1>Raw HTML slide</h1>"]}]}
    llm.inject("slide_planner", json.dumps(html))
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    failed = [e for e in container.task_service.events(task.task_id)
              if e.type == "agent.validation_failed" and e.agent_id == "slide_planner"]
    assert "content_blocks" in failed[0].data["error"]


async def test_render_failure_fails_the_task_and_keeps_it_inspectable(make_container) -> None:
    container = make_container(presentation_renderer=FailingRenderer())
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED
    error = task.errors[-1]
    assert error.node_id == "render_presentation" and "disk full" in error.message
    saved = container.task_service.get(task.task_id)  # persisted and inspectable
    assert saved.status == TaskStatus.FAILED and saved.errors[-1].message == error.message
    states = saved.workflow.node_states
    assert states["build_presentation"].status.value == "COMPLETED"
    assert states["render_presentation"].status.value == "FAILED" and "disk full" in states["render_presentation"].error
    Presentation.model_validate(states["build_presentation"].output)  # the built presentation is still there
    names = set(arts_of(container, task))
    assert "slide_plan" in names and "presentation" not in names
    assert not any(a.type == ArtifactType.PRESENTATION for a in container.task_service.artifacts(task.task_id))
    failed = [e for e in presentation_events(container, task) if e.type == "presentation.failed"]
    assert failed[-1].data["stage"] == "render" and "disk full" in failed[-1].data["error"]


async def test_rejected_lessons_get_no_presentation(make_container) -> None:
    llm = MockLLMProvider({**default_responders(), "content_reviewer": always_reject})
    container = make_container(llm=llm, max_revisions=1)
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED and task.errors[-1].node_id == "teach_review"
    assert llm.calls["slide_planner"] == 0
    assert task.workflow.node_states["slide_plan"].status.value == "PENDING"
    assert not presentation_events(container, task)


async def test_lesson_accepted_with_warnings_gets_no_presentation(make_container) -> None:
    llm = MockLLMProvider({**default_responders(), "content_reviewer": always_reject})
    container = make_container(llm=llm, max_revisions=1, revision_exhausted_policy="accept_with_warnings")
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    for node in ("slide_plan", "validate_slide_plan", "build_presentation", "render_presentation"):
        assert task.workflow.node_states[node].status.value == "SKIPPED", node
    assert llm.calls["slide_planner"] == 0
    assert {"slide_plan", "presentation"}.isdisjoint(arts_of(container, task))
    assert any("presentations are only generated for an approved lesson" in w for w in task.result.warnings)
    assert task.workflow.node_states["update_learner"].status.value == "COMPLETED"


async def test_missing_optional_visual_is_not_replaced(make_container) -> None:
    # The optional section-2 illustration cannot be found or generated (the existing visual policy continues).
    container = make_container(image_search_provider=SearchWithout("match"),
                               image_generation_provider=GenerationFailing(prefix="Illustration"))
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    images = {a.artifact_id for a in container.task_service.artifacts(task.task_id)
              if a.type == ArtifactType.IMAGE_ASSET}
    assert len(images) == 3
    deck = deck_of(container, task)
    assert set(deck.image_artifact_ids()) == images  # only existing assets, nothing invented in its place
    doc = pptx_of(container, task)
    assert sum(1 for page in doc.slides for s in page.shapes if s.shape_type == PICTURE) == 3
    assert any("v2_illustration" in w for w in task.result.warnings)


async def test_crash_after_slide_planning_resumes_without_rerunning_earlier_work(make_container, tmp_path) -> None:
    data_dir = tmp_path / "shared"
    first = make_container(observers=[crash_after("slide_plan")], data_dir=data_dir)
    learner_id = add_demo_learner(first)
    task = await first.task_service.create_and_run(request="Create an A2 lesson about football.",
                                                   learner_id=learner_id, user_id="u1")
    task = await first.task_service.submit_assessment(task.task_id, answers_for(task))
    with pytest.raises(SimulatedCrash):
        await first.task_service.submit_assessment(task.task_id, answers_for(task))
    crashed = first.task_service.get(task.task_id)
    assert crashed.workflow.node_states["slide_plan"].status.value == "COMPLETED"
    assert crashed.workflow.node_states["validate_slide_plan"].status.value == "PENDING"
    before = {a.artifact_id for a in first.task_service.artifacts(task.task_id)}

    llm = MockLLMProvider(default_responders())
    generation = GenerationFailing(prefix="")  # any image generation now would fail the task
    second = make_container(llm=llm, data_dir=data_dir, image_generation_provider=generation)
    resumed = await second.task_service.resume(task.task_id)
    assert resumed.status == TaskStatus.COMPLETED, resumed.errors
    assert sum(llm.calls.values()) == 0 and generation.calls == 0  # no diagnostic, research, teaching or images
    events = second.task_service.events(task.task_id)
    started = [e.node_id for e in events if e.type == "node.started"]
    for node in ("diagnose_1", "research", "plan", "teach_review", "visual", "store_artifacts", "slide_plan"):
        assert started.count(node) == 1, node
    for node in ("validate_slide_plan", "store_slide_plan", "build_presentation", "render_presentation"):
        assert started.count(node) == 1, node
    after = {a.artifact_id: a for a in second.task_service.artifacts(task.task_id)}
    assert before <= set(after)
    assert {a.name for a in after.values() if a.artifact_id not in before} == {"slide_plan", "presentation"}


async def test_rerunning_the_render_stage_reuses_the_presentation_artifact(make_container, tmp_path) -> None:
    data_dir = tmp_path / "shared"
    first = make_container(data_dir=data_dir)

    class Stop(BaseException):
        pass

    def stop_once(event):  # the process dies after storing the artifact, before the node is checkpointed
        if event.type == "presentation.artifact_created":
            raise Stop()

    first.events.subscribe(stop_once)
    learner_id = add_demo_learner(first)
    task = await first.task_service.create_and_run(request="Create an A2 lesson about football.",
                                                   learner_id=learner_id, user_id="u1")
    task = await first.task_service.submit_assessment(task.task_id, answers_for(task))
    with pytest.raises(Stop):
        await first.task_service.submit_assessment(task.task_id, answers_for(task))
    stored = first.artifacts.find(task.task_id, "presentation")
    assert stored is not None
    assert first.task_service.get(task.task_id).workflow.node_states["render_presentation"].status.value != "COMPLETED"

    second = make_container(data_dir=data_dir)
    resumed = await second.task_service.resume(task.task_id)
    assert resumed.status == TaskStatus.COMPLETED, resumed.errors
    presentations = [a for a in second.task_service.artifacts(task.task_id) if a.type == ArtifactType.PRESENTATION]
    assert [(a.artifact_id, a.version) for a in presentations] == [(stored.artifact_id, 1)]  # no duplicate
    created = [e for e in second.task_service.events(task.task_id) if e.type == "presentation.artifact_created"]
    assert [e.data["reused"] for e in created] == [False, True]
    assert sum(1 for e in second.task_service.events(task.task_id)
               if e.type == "artifact.created" and e.data["artifact_type"] == "PRESENTATION") == 1
    objects = list((data_dir / "objects" / "objects").rglob("*.pptx"))
    assert [p.name for p in objects] == [f"{stored.content_hash}.pptx"]


async def test_identical_lessons_render_byte_identical_files(make_container) -> None:
    a, b = make_container(), make_container()
    first, second = await run_lesson(a), await run_lesson(b)
    pa, pb = arts_of(a, first)["presentation"], arts_of(b, second)["presentation"]
    assert pa.content_hash == pb.content_hash  # deterministic ids and a reproducible renderer


async def test_mock_renderer_and_4_3_settings(make_container) -> None:
    container = make_container(presentation_renderer="mock")
    task = await run_lesson(container)
    art = arts_of(container, task)["presentation"]
    assert art.media_type == "application/json" and art.metadata["renderer"] == "mock-presentation"
    doc = json.loads(container.artifacts.read(art.artifact_id))
    assert [s["slide_id"] for s in doc["slides"]] == [s.slide_id for s in deck_of(container, task).slides]

    narrow = make_container(presentation_aspect_ratio="4:3")
    task = await run_lesson(narrow)
    doc = pptx_of(narrow, task)
    assert doc.slide_width * 3 == doc.slide_height * 4
