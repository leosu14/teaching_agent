"""The audio stages inside the lesson workflow: audio planning, the plan validation gate, TTS, audio validation,
AUDIO_ASSET artifacts, the presentation timeline, gating on the review and the presentation, the failure policy,
events, cost, language and voice, resume and idempotency, and the PRESENTATION -> timeline -> AUDIO_ASSET links."""

from __future__ import annotations

import hashlib
import io
import json
import wave
import zipfile

import pytest
from pptx import Presentation as PptxDocument

from app.config.settings import Settings
from app.providers.core.errors import ProviderError
from app.providers.llm.mock import MockLLMProvider
from app.providers.llm.mock_responders import default_responders
from app.providers.tts.mock import MockTTSProvider
from app.schemas.artifact import ArtifactType
from app.schemas.audio import (
    AudioAssetMetadata,
    AudioPlan,
    AudioPlanningInput,
    NarrationResult,
    PresentationTimeline,
)
from app.schemas.presentation import SlideDeckPlan
from app.schemas.task import TaskStatus
from app.services.container import build_container
from tests.conftest import add_demo_learner, answers_for, run_lesson
from tests.integration.test_presentation_flow import FailingRenderer
from tests.integration.test_resume import SimulatedCrash, crash_after
from tests.integration.test_review_loop import always_reject
from tests.integration.test_visual_flow import GenerationFailing

AUDIO_NODES = ["audio_plan", "validate_audio_plan", "store_audio_plan", "synthesize_audio", "audio_policy",
               "audio_timeline"]
PROVIDERS = ("image_generation_provider", "presentation_renderer", "tts_provider")


class SelectiveTTS(MockTTSProvider):
    """The mock provider, except for texts containing `marker`: those fail, or return audio that lies about itself."""

    def __init__(self, marker: str, mode: str = "fail") -> None:
        super().__init__()
        self.marker, self.mode = marker, mode

    async def synthesize(self, request):
        speech = await super().synthesize(request)
        if self.marker not in request.text:
            return speech
        if self.mode == "fail":
            raise ProviderError("voice service rejected the text", transient=False)
        if self.mode == "fake":  # arbitrary bytes posing as a WAV file
            return speech.model_copy(update={"content": b"RIFF....not really audio" * 20})
        return speech.model_copy(update={"duration": speech.duration + 3.0})  # declared duration is wrong


@pytest.fixture
def make_container(tmp_path):
    made = []

    def make(*, llm: MockLLMProvider | None = None, observers=(), data_dir=None, **kwargs):
        providers = {k: kwargs.pop(k) for k in PROVIDERS if k in kwargs and not isinstance(kwargs[k], str)}
        c = build_container(Settings(data_dir=data_dir or tmp_path / f"d{len(made)}", log_json=False, **kwargs),
                            llm_providers={"mock": llm or MockLLMProvider(default_responders())},
                            observers=observers, **providers)
        made.append(c)
        return c

    yield make
    for c in made:
        c.close()


def arts_of(container, task) -> dict:
    return {a.name: a for a in container.task_service.artifacts(task.task_id)}


def json_of(container, artifact, model):
    return model.model_validate_json(container.artifacts.read(artifact.artifact_id))


def audio_events(container, task) -> list:
    return [e for e in container.task_service.events(task.task_id)
            if e.type.startswith(("audio_planning.", "audio_plan.", "tts.", "audio.", "timeline."))]


def states(task, *nodes) -> list[str]:
    return [task.workflow.node_states[n].status.value for n in nodes]


async def test_approved_lesson_is_narrated_with_real_wav_assets_and_a_timeline(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    tts = MockTTSProvider()
    container = make_container(llm=llm, tts_provider=tts)
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    order = [n for n in task.workflow.execution_order if n in {"render_presentation", *AUDIO_NODES, "update_learner"}]
    assert order == ["render_presentation", *AUDIO_NODES, "update_learner"]

    arts = arts_of(container, task)
    deck = json_of(container, arts["slide_plan"], SlideDeckPlan)
    plan = json_of(container, arts["audio_plan"], AudioPlan)
    pres = arts["presentation"]

    # The planner saw the slides' speakable text only: no image, table or citation blocks.
    planning = AudioPlanningInput.model_validate(next(r for r in llm.requests if r.agent_id == "audio_planner").input_payload)
    assert [s.slide_id for s in planning.slides] == [s.slide_id for s in deck.slides]
    payload = json.dumps(planning.model_dump(mode="json"))
    assert "art_" not in payload and "file://" not in payload and "cite_" not in payload

    # AudioPlan: slide-mapped segments in deck order, the lesson's language, a catalog voice, every source type.
    assert plan.task_id == task.task_id and plan.deck_id == deck.deck_id
    assert plan.language == "en" and plan.voice == "mock-en-US-1" and all(s.voice == plan.voice for s in plan.segments)
    position = {s.slide_id: i for i, s in enumerate(deck.slides)}
    assert [position[s.slide_id] for s in plan.segments] == sorted(position[s.slide_id] for s in plan.segments)
    assert {s.source_type.value for s in plan.segments} == {
        "slide_title", "slide_content", "speaker_notes", "exercise_instructions", "answer_explanation"}
    references = next(s for s in deck.slides if s.slide_type.value == "references")
    assert not plan.for_slide(references.slide_id)  # nothing on a references slide is read aloud
    research = json.loads(container.artifacts.read(arts["research_bundle"].artifact_id))
    spoken = " ".join(s.text for s in plan.segments)
    assert not any(c["citation_id"] in spoken or c["url"] in spoken for c in research["citations"])
    assert all(s.start_time is None for s in plan.segments)  # the stored plan is the planned one

    # One AUDIO_ASSET per segment: a real WAV, measured metadata, content-addressed, linked to the plan.
    assets = {a.metadata["segment_id"]: a for a in arts.values() if a.type == ArtifactType.AUDIO_ASSET}
    assert set(assets) == {s.segment_id for s in plan.segments} and tts.calls == len(plan.segments)
    for seg in plan.segments:
        art = assets[seg.segment_id]
        meta = AudioAssetMetadata.model_validate(art.metadata)
        data = container.artifacts.read(art.artifact_id)
        with wave.open(io.BytesIO(data)) as wav:
            assert wav.getframerate() == meta.sample_rate == 16000 and wav.getnchannels() == meta.channels == 1
            assert wav.getnframes() / wav.getframerate() == pytest.approx(meta.duration)
        assert hashlib.sha256(data).hexdigest() == art.content_hash == meta.checksum
        assert art.media_type == meta.media_type == "audio/wav" and meta.format == "wav"
        assert (meta.language, meta.voice, meta.provider, meta.model) == ("en", plan.voice, "mock", "mock-tts-1")
        assert meta.slide_id == seg.slide_id and meta.text == seg.text and meta.validation.valid
        assert meta.object_key == f"objects/sha256/{meta.checksum[:2]}/{meta.checksum}.wav"
        assert art.parent_ids == [arts["audio_plan"].artifact_id]
    assert arts["audio_plan"].type == ArtifactType.AUDIO_PLAN
    assert arts["audio_plan"].parent_ids == [pres.artifact_id, arts["slide_plan"].artifact_id,
                                             arts["lesson"].artifact_id]

    # The timeline: every slide in order, contiguous, built from the measured durations.
    tl_art = arts["presentation_timeline"]
    timeline = json_of(container, tl_art, PresentationTimeline)
    assert [s.slide_id for s in timeline.slides] == [s.slide_id for s in deck.slides]
    for seg_t in timeline.segments:
        assert seg_t.duration == pytest.approx(assets[seg_t.segment_id].metadata["duration"], abs=0.0006)
        assert seg_t.audio_artifact_id == assets[seg_t.segment_id].artifact_id
    estimates = {s.segment_id: s.expected_duration for s in plan.segments}
    assert any(abs(t.duration - estimates[t.segment_id]) > 0.1 for t in timeline.segments)  # estimates were not used
    measured = sum(s.duration for s in timeline.segments)
    assert timeline.slide(references.slide_id).audio_segment_refs == [] and timeline.slide(
        references.slide_id).duration == 3.0
    assert timeline.duration == timeline.slides[-1].end_time > measured
    assert tl_art.type == ArtifactType.PRESENTATION_TIMELINE
    assert tl_art.parent_ids == [pres.artifact_id, arts["audio_plan"].artifact_id, *timeline.audio_artifact_ids()]
    timed = AudioPlan.model_validate(task.workflow.node_states["audio_timeline"].output["plan"])
    assert all(s.start_time is not None and s.end_time > s.start_time for s in timed.segments)

    # Task result and lineage.
    result_ids = {a.artifact_id for a in task.result.artifacts}
    assert {arts["audio_plan"].artifact_id, tl_art.artifact_id, *(a.artifact_id for a in assets.values())} <= result_ids
    assert tl_art.artifact_id in task.artifact_ids
    assert {"presentation", "audio_plan", "slide_plan", "lesson", "research_bundle"} <= {
        a.name for a in container.artifacts.lineage(tl_art.artifact_id)}
    assert task.result.warnings == []


async def test_timeline_links_presentation_slides_and_audio_without_touching_the_pptx(make_container) -> None:
    container = make_container()
    task = await run_lesson(container)
    arts = arts_of(container, task)
    pres = arts["presentation"]
    pptx_bytes = container.artifacts.read(pres.artifact_id)
    timeline = json_of(container, arts["presentation_timeline"], PresentationTimeline)
    deck = json_of(container, arts["slide_plan"], SlideDeckPlan)

    assert timeline.presentation_artifact_id == pres.artifact_id and timeline.deck_id == deck.deck_id
    assert pres.metadata["deck_id"] == deck.deck_id
    doc = PptxDocument(io.BytesIO(pptx_bytes))
    assert len(doc.slides) == len(timeline.slides) == pres.metadata["slides"]
    assert [s.shapes.title.text for s in doc.slides] == [s.title for s in deck.slides]
    assert [s.slide_id for s in timeline.slides] == [s.slide_id for s in deck.slides]
    for aid in timeline.audio_artifact_ids():
        assert container.artifacts.get(aid).type == ArtifactType.AUDIO_ASSET
    # The PPTX is unchanged: one version, its bytes match its checksum, and it carries no audio.
    presentations = [a for a in container.task_service.artifacts(task.task_id) if a.type == ArtifactType.PRESENTATION]
    assert [a.version for a in presentations] == [1]
    assert hashlib.sha256(pptx_bytes).hexdigest() == pres.content_hash
    assert not [n for n in zipfile.ZipFile(io.BytesIO(pptx_bytes)).namelist() if n.endswith((".wav", ".mp3"))]


async def test_audio_events_and_cost(make_container) -> None:
    container = make_container()
    task = await run_lesson(container)
    events = audio_events(container, task)
    types = [e.type for e in events]
    assert types[:3] == ["audio_planning.started", "audio_plan.created", "audio_plan.validated"]
    assert types[-2:] == ["timeline.created", "audio.completed"]
    voiced = types[3:-2]
    assert voiced == ["tts.started", "tts.completed", "audio.asset_created"] * (len(voiced) // 3)
    plan = json_of(container, arts_of(container, task)["audio_plan"], AudioPlan)
    assert len(voiced) // 3 == len(plan.segments)
    assert all(e.data["ok"] for e in events if e.type == "tts.completed")
    assert all(e.task_id == task.task_id for e in events)
    assert {e.node_id for e in events if e.type.startswith("tts.")} == {"synthesize_audio"}
    assert events[-1].data["status"] == "complete" and events[-1].data["assets"] == len(plan.segments)

    cost = task.cost
    planner = cost.by_agent["audio_planner"]
    assert planner.calls == 1 and planner.cost_usd > 0 and planner.usage.total_tokens > 0
    line = cost.by_service["tts:mock"]
    assert line.calls == len(plan.segments) and line.cost_usd == 0.0 and line.estimated_cost_usd is None
    assert line.units["characters"] == sum(len(s.text) for s in plan.segments)
    timeline = json_of(container, arts_of(container, task)["presentation_timeline"], PresentationTimeline)
    assert line.units["seconds"] == pytest.approx(sum(s.duration for s in timeline.segments), abs=0.01)
    assert "tts:mock" not in cost.by_model  # speech is a service, never mixed into model usage


async def test_rejected_lessons_and_lessons_without_a_presentation_get_no_audio(make_container) -> None:
    llm = MockLLMProvider({**default_responders(), "content_reviewer": always_reject})
    tts = MockTTSProvider()
    container = make_container(llm=llm, tts_provider=tts, max_revisions=1)
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED and task.errors[-1].node_id == "teach_review"
    assert llm.calls["audio_planner"] == 0 and tts.calls == 0 and not audio_events(container, task)

    llm = MockLLMProvider({**default_responders(), "content_reviewer": always_reject})
    container = make_container(llm=llm, tts_provider=tts, max_revisions=1,
                               revision_exhausted_policy="accept_with_warnings")
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    assert states(task, *AUDIO_NODES) == ["SKIPPED"] * len(AUDIO_NODES)  # no visuals, no presentation, no audio
    assert llm.calls["audio_planner"] == 0 and tts.calls == 0
    assert any("audio is only generated for the presentation of an approved lesson" in w for w in task.result.warnings)
    assert task.workflow.node_states["update_learner"].status.value == "COMPLETED"


async def test_presentation_failure_means_no_audio(make_container) -> None:
    tts = MockTTSProvider()
    container = make_container(presentation_renderer=FailingRenderer(), tts_provider=tts)
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED and task.errors[-1].node_id == "render_presentation"
    assert states(task, *AUDIO_NODES) == ["PENDING"] * len(AUDIO_NODES)
    assert tts.calls == 0 and not audio_events(container, task)
    assert not [a for a in container.task_service.artifacts(task.task_id) if a.type.value.startswith("AUDIO")]


async def test_required_narration_failure_fails_the_task_and_keeps_generated_audio(make_container) -> None:
    tts = SelectiveTTS(marker="Practise:")  # the objectives slide's narration (required)
    container = make_container(tts_provider=tts)
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED
    error = task.errors[-1]
    assert error.node_id == "audio_policy" and "required narration failed" in error.message
    assert "voice service rejected the text" in error.message
    narration = NarrationResult.model_validate(task.workflow.node_states["synthesize_audio"].output)
    assert narration.status == "failed" and [f.stage for f in narration.failures] == ["tts"]
    failed_id = narration.failures[0].segment_id
    assets = {a.metadata["segment_id"] for a in container.task_service.artifacts(task.task_id)
              if a.type == ArtifactType.AUDIO_ASSET}
    plan = json_of(container, arts_of(container, task)["audio_plan"], AudioPlan)
    assert assets == {s.segment_id for s in plan.segments} - {failed_id}  # every other segment was kept
    events = audio_events(container, task)
    assert any(e.type == "tts.completed" and not e.data["ok"] and e.data["segment_id"] == failed_id for e in events)
    failed = [e for e in events if e.type == "audio.failed"]
    assert failed and failed[-1].data["stage"] == "tts"
    assert "presentation_timeline" not in arts_of(container, task)


async def test_required_failure_continues_when_the_policy_says_so(make_container) -> None:
    container = make_container(tts_provider=SelectiveTTS(marker="Practise:"), audio_failure_policy="continue")
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    timeline = json_of(container, arts_of(container, task)["presentation_timeline"], PresentationTimeline)
    assert len(timeline.missing_segments) == 1
    assert any("audio policy is 'continue'" in w and timeline.missing_segments[0] in w for w in task.result.warnings)


async def test_optional_narration_failure_continues_with_a_warning(make_container) -> None:
    tts = SelectiveTTS(marker="Compare with the lesson examples")  # answer explanations are optional narration
    container = make_container(tts_provider=tts)
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    plan = json_of(container, arts_of(container, task)["audio_plan"], AudioPlan)
    optional = [s for s in plan.segments if not s.required]
    assert optional and all(s.source_type.value == "answer_explanation" for s in optional)
    timeline = json_of(container, arts_of(container, task)["presentation_timeline"], PresentationTimeline)
    assert sorted(timeline.missing_segments) == sorted(s.segment_id for s in optional)
    answer_slide = optional[0].slide_id
    assert timeline.slide(answer_slide).audio_segment_refs == [] and timeline.slide(answer_slide).duration > 0
    warnings = [w for w in task.result.warnings if w.startswith("Optional narration")]
    assert len(warnings) == len(optional)  # not skipped silently
    assets = [a for a in container.task_service.artifacts(task.task_id) if a.type == ArtifactType.AUDIO_ASSET]
    assert len(assets) == len(plan.segments) - len(optional)


@pytest.mark.parametrize("mode", ["fake", "lie"])
async def test_invalid_audio_is_rejected_before_it_becomes_an_asset(make_container, mode) -> None:
    container = make_container(tts_provider=SelectiveTTS(marker="Practise:", mode=mode))
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED and task.errors[-1].node_id == "audio_policy"
    narration = NarrationResult.model_validate(task.workflow.node_states["synthesize_audio"].output)
    failure = narration.failures[0]
    assert failure.stage == "validate"
    assert {e.code for e in failure.errors} == ({"unreadable_audio"} if mode == "fake" else {"duration_mismatch"})
    rejected = [e for e in audio_events(container, task) if e.type == "audio.validation_failed"]
    assert [e.data["segment_id"] for e in rejected] == [failure.segment_id]
    assert failure.segment_id not in {a.metadata["segment_id"] for a in container.task_service.artifacts(task.task_id)
                                      if a.type == ArtifactType.AUDIO_ASSET}


async def test_language_and_voice_settings(make_container) -> None:
    container = make_container(audio_language="es-ES", audio_voice="mock-es-ES-2")
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    plan = json_of(container, arts_of(container, task)["audio_plan"], AudioPlan)
    assert plan.language == "es-ES" and plan.voice == "mock-es-ES-2"
    assert {(s.language, s.voice) for s in plan.segments} == {("es-ES", "mock-es-ES-2")}
    meta = [a.metadata for a in container.task_service.artifacts(task.task_id) if a.type == ArtifactType.AUDIO_ASSET]
    assert {(m["language"], m["voice"]) for m in meta} == {("es-ES", "mock-es-ES-2")}

    zh = make_container(audio_language="zh-CN")
    task = await run_lesson(zh)
    assert task.status == TaskStatus.COMPLETED, task.errors
    plan = json_of(zh, arts_of(zh, task)["audio_plan"], AudioPlan)
    assert (plan.language, plan.voice) == ("zh-CN", "mock-zh-CN-1")


@pytest.mark.parametrize("settings, message", [
    ({"audio_language": "sw-KE"}, "no voice for sw-KE"),
    ({"audio_language": "es-ES", "audio_voice": "mock-en-US-1"}, "does not speak es-ES"),
])
async def test_unsupported_language_or_voice_fails_explicitly(make_container, settings, message) -> None:
    tts = MockTTSProvider()
    container = make_container(tts_provider=tts, **settings)
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED and task.errors[-1].node_id == "audio_plan"
    assert message in task.errors[-1].message and tts.calls == 0
    failed = [e for e in audio_events(container, task) if e.type == "audio.failed"]
    assert failed and failed[-1].data["stage"] == "planning"


async def test_planner_corrects_a_plan_that_fails_validation(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    container = make_container(llm=llm)
    first = {"segments": [
        {"slide_id": "s01", "source_type": "slide_title", "source_ref": "s01.title", "text": "Football"},
        {"slide_id": "s99", "source_type": "speaker_notes", "source_ref": "s99.notes", "text": "Invented slide."},
        {"slide_id": "s02", "source_type": "slide_content", "source_ref": "s02.content",
         "text": "Read https://example.org/source aloud."}]}
    llm.inject("audio_planner", json.dumps(first))
    task = await run_lesson(container)
    assert task.status == TaskStatus.COMPLETED, task.errors
    assert llm.calls["audio_planner"] == 2
    second = AudioPlanningInput.model_validate([r for r in llm.requests if r.agent_id == "audio_planner"][1].input_payload)
    assert any("unknown_slide" in c and "s99" in c for c in second.corrections)
    assert any("spoken_metadata" in c for c in second.corrections)
    failed = [e for e in container.task_service.events(task.task_id)
              if e.type == "agent.validation_failed" and e.agent_id == "audio_planner"]
    assert len(failed) == 1 and failed[0].data["stage"] == "audio_plan"


async def test_a_plan_that_stays_invalid_fails_before_any_speech(make_container) -> None:
    llm = MockLLMProvider(default_responders())
    tts = MockTTSProvider()
    container = make_container(llm=llm, tts_provider=tts)
    broken = {"segments": [{"slide_id": "s99", "source_type": "speaker_notes", "source_ref": "x", "text": "Hi."}]}
    llm.inject("audio_planner", *[json.dumps(broken)] * 3)
    task = await run_lesson(container)
    assert task.status == TaskStatus.FAILED and task.errors[-1].node_id == "validate_audio_plan"
    assert "unknown_slide" in task.errors[-1].message and tts.calls == 0
    assert "audio_plan" not in arts_of(container, task)
    failed = [e for e in audio_events(container, task) if e.type == "audio.failed"]
    assert failed[-1].data["stage"] == "validation"


async def start_until_crash(container, crash_type=SimulatedCrash):
    learner_id = add_demo_learner(container)
    task = await container.task_service.create_and_run(request="Create an A2 lesson about football.",
                                                       learner_id=learner_id, user_id="u1")
    task = await container.task_service.submit_assessment(task.task_id, answers_for(task))
    with pytest.raises(crash_type):
        await container.task_service.submit_assessment(task.task_id, answers_for(task))
    return task


async def test_crash_after_audio_planning_resumes_from_the_audio_plan(make_container, tmp_path) -> None:
    data_dir = tmp_path / "shared"
    first = make_container(observers=[crash_after("audio_plan")], data_dir=data_dir)
    task = await start_until_crash(first)
    crashed = first.task_service.get(task.task_id)
    assert states(crashed, "audio_plan", "validate_audio_plan") == ["COMPLETED", "PENDING"]
    planned = AudioPlan.model_validate(crashed.workflow.node_states["audio_plan"].output)

    llm = MockLLMProvider(default_responders())
    generation = GenerationFailing(prefix="")  # any image generation now would fail the task
    tts = MockTTSProvider()
    second = make_container(llm=llm, data_dir=data_dir, image_generation_provider=generation, tts_provider=tts)
    resumed = await second.task_service.resume(task.task_id)
    assert resumed.status == TaskStatus.COMPLETED, resumed.errors
    assert sum(llm.calls.values()) == 0 and generation.calls == 0  # no research, teaching, images or re-planning
    assert tts.calls == len(planned.segments)
    started = [e.node_id for e in second.task_service.events(task.task_id) if e.type == "node.started"]
    for node in ("research", "teach_review", "visual", "slide_plan", "render_presentation", "audio_plan",
                 *AUDIO_NODES[1:]):
        assert started.count(node) == 1, node
    stored = json_of(second, arts_of(second, resumed)["audio_plan"], AudioPlan)
    assert stored == planned
    presentations = [a for a in second.task_service.artifacts(task.task_id) if a.type == ArtifactType.PRESENTATION]
    assert len(presentations) == 1  # the PPTX was not regenerated


class Stop(BaseException):
    pass


@pytest.mark.parametrize("stop_after", [3, "last"])
async def test_rerunning_the_narration_reuses_existing_audio_assets(make_container, tmp_path, stop_after) -> None:
    """The process dies inside the TTS stage after some (or all) assets were stored. Running the stage again reuses
    every stored asset without synthesizing it again and creates no duplicate AUDIO_ASSET."""
    data_dir = tmp_path / "shared"
    first_tts = MockTTSProvider()
    first = make_container(data_dir=data_dir, tts_provider=first_tts)
    seen = []

    def stop(event):
        if event.type == "audio.asset_created":
            seen.append(event)
            plan_size = len(first.task_service.get(event.task_id).workflow.node_states["audio_plan"].output["segments"])
            if len(seen) == (plan_size if stop_after == "last" else stop_after):
                raise Stop()

    first.events.subscribe(stop)
    task = await start_until_crash(first, Stop)
    before = {a.artifact_id: a for a in first.task_service.artifacts(task.task_id) if a.type == ArtifactType.AUDIO_ASSET}
    assert len(before) == len(seen) and first_tts.calls == len(seen)
    assert first.task_service.get(task.task_id).workflow.node_states["synthesize_audio"].status.value != "COMPLETED"

    tts = MockTTSProvider()
    second = make_container(data_dir=data_dir, tts_provider=tts)
    resumed = await second.task_service.resume(task.task_id)
    assert resumed.status == TaskStatus.COMPLETED, resumed.errors
    plan = json_of(second, arts_of(second, resumed)["audio_plan"], AudioPlan)
    assert tts.calls == len(plan.segments) - len(seen)  # only the missing segments were synthesized
    after = [a for a in second.task_service.artifacts(task.task_id) if a.type == ArtifactType.AUDIO_ASSET]
    assert len(after) == len(plan.segments) and all(a.version == 1 for a in after)  # no duplicates
    assert set(before) <= {a.artifact_id for a in after}
    created = [e for e in second.task_service.events(task.task_id) if e.type == "audio.asset_created"]
    assert sum(e.data["reused"] for e in created) == len(seen)
    narration = NarrationResult.model_validate(resumed.workflow.node_states["synthesize_audio"].output)
    assert sum(a.reused for a in narration.assets) == len(seen)
    wavs = list((data_dir / "objects" / "objects").rglob("*.wav"))
    assert len(wavs) == len({a.content_hash for a in after})


async def test_rerunning_the_timeline_reuses_the_timeline_artifact(make_container, tmp_path) -> None:
    data_dir = tmp_path / "shared"
    first = make_container(data_dir=data_dir)

    def stop(event):  # dies after the timeline is stored, before its node is checkpointed
        if event.type == "audio.completed":
            raise Stop()

    first.events.subscribe(stop)
    task = await start_until_crash(first, Stop)
    stored = first.artifacts.find(task.task_id, "presentation_timeline")
    assert stored is not None

    tts = MockTTSProvider()
    second = make_container(data_dir=data_dir, tts_provider=tts)
    resumed = await second.task_service.resume(task.task_id)
    assert resumed.status == TaskStatus.COMPLETED, resumed.errors
    assert tts.calls == 0
    timelines = [a for a in second.task_service.artifacts(task.task_id) if a.type == ArtifactType.PRESENTATION_TIMELINE]
    assert [(a.artifact_id, a.version) for a in timelines] == [(stored.artifact_id, 1)]


async def test_identical_lessons_produce_identical_audio_stored_once(make_container, tmp_path) -> None:
    container = make_container()
    first, second = await run_lesson(container), await run_lesson(container)
    assert first.status == second.status == TaskStatus.COMPLETED
    hashes = [sorted(a.content_hash for a in container.task_service.artifacts(t.task_id)
                     if a.type == ArtifactType.AUDIO_ASSET) for t in (first, second)]
    assert hashes[0] == hashes[1]  # deterministic plan, voice and provider
    reused = [e.data["reused_object"] for e in container.task_service.events(second.task_id)
              if e.type == "tts.completed"]
    assert reused and all(reused)  # the second task's bytes were already in the store
    wavs = list((tmp_path / "d0" / "objects" / "objects").rglob("*.wav"))
    assert len(wavs) == len(set(hashes[0]))
