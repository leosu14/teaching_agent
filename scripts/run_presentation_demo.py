"""End-to-end demo of the Presentation slice.

    python scripts/run_presentation_demo.py [--out lesson.pptx]

Runs the same lesson as run_demo.py offline with the mock providers (Diagnostic -> Research -> Planner -> Teacher ->
Reviewer -> Visual -> Slide Planning -> Presentation Build -> Presentation Render), then walks through the
presentation stages: the completed lesson, its ResearchBundle and IMAGE_ASSET artifacts, the SlideDeckPlan and its
validation, the built presentation, the rendered PPTX and the stored PRESENTATION artifact. Finally it opens the
.pptx with python-pptx to show it is a real PowerPoint file with the images placed. It uses the same TaskService
as the API.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import sys
import tempfile
from pathlib import Path

from pptx import Presentation as PptxDocument

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from app.config.settings import Settings  # noqa: E402
from app.observability.logging import configure_logging  # noqa: E402
from app.schemas.artifact import ArtifactType  # noqa: E402
from app.schemas.lesson import LessonContent  # noqa: E402
from app.schemas.presentation import (  # noqa: E402
    Presentation,
    PresentationArtifactMetadata,
    SlideDeckPlan,
    SlidePlanValidationReport,
)
from app.schemas.research import ResearchBundle  # noqa: E402
from app.schemas.task import Task, TaskStatus  # noqa: E402
from app.schemas.visual import ImageAsset  # noqa: E402
from app.services.container import Container, build_container  # noqa: E402
from scripts.run_demo import run_demo  # noqa: E402

PICTURE = 13  # MSO_SHAPE_TYPE.PICTURE


def _clip(text: str, n: int = 80) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 3] + "..."


def _heading(out, title: str) -> None:
    out("\n" + "=" * 72 + f"\n{title}\n" + "=" * 72)


def report_presentation(container: Container, task: Task, out=print, save_to: Path | None = None) -> None:
    arts = {a.name: a for a in container.task_service.artifacts(task.task_id)}
    nodes = task.workflow.node_states
    name_of = lambda aid: container.artifacts.get(aid).name  # noqa: E731

    _heading(out, "1. Completed lesson")
    lesson = LessonContent.model_validate_json(container.artifacts.read(arts["lesson"].artifact_id))
    out(f"{lesson.title} ({lesson.level}); review: {task.result.review_verdict} after {task.result.revisions} "
        f"revision(s); lesson artifact {arts['lesson'].artifact_id}")
    for s in lesson.sections:
        out(f"  {s.section_id}: {s.heading}  citations={','.join(s.citations)}  visuals={len(s.visuals)}")

    _heading(out, "2. ResearchBundle")
    research = ResearchBundle.model_validate_json(container.artifacts.read(arts["research_bundle"].artifact_id))
    out(f"{research.research_id}: status={research.status} sources={len(research.sources)} "
        f"citations={len(research.citations)}")

    _heading(out, "3. IMAGE_ASSET artifacts")
    images = [a for a in arts.values() if a.type == ArtifactType.IMAGE_ASSET]
    for art in sorted(images, key=lambda a: a.name):
        asset = ImageAsset.model_validate(art.metadata)
        out(f"  {art.artifact_id} {art.name:<24} {asset.origin:<9} {asset.width}x{asset.height} "
            f"section={asset.lesson_section_id} sha256={asset.object.checksum[:12]}...")

    _heading(out, "4. SlideDeckPlan (slide planner)")
    plan_art = arts["slide_plan"]
    deck = SlideDeckPlan.model_validate_json(container.artifacts.read(plan_art.artifact_id))
    out(f"{deck.deck_id}: {deck.title!r} language={deck.language} level={deck.level} topic={deck.topic}")
    out(f"objective: {_clip(deck.objective, 90)}")
    out(f"artifact: {plan_art.artifact_id} {plan_art.type.value} v{plan_art.version} "
        f"parent: {', '.join(name_of(p) for p in plan_art.parent_ids)}")
    for s in deck.slides:
        blocks = ", ".join(b.kind for b in s.content_blocks) or "-"
        out(f"  {s.order:>2}. {s.slide_type.value:<12} {s.layout.value:<13} {_clip(s.title, 40)!r}")
        out(f"      blocks: {blocks}; images: {', '.join(s.visual_refs) or '-'}; "
            f"citations: {', '.join(s.citation_refs) or '-'}; sections: {', '.join(s.section_refs) or '-'}")

    _heading(out, "5. Validation")
    report = SlidePlanValidationReport.model_validate(nodes["validate_slide_plan"].output)
    out(f"valid={report.valid} by {report.validator}; errors={len(report.errors)}")
    out(f"checks: {', '.join(report.checks)}")

    _heading(out, "6. Presentation build (renderer-independent)")
    built = Presentation.model_validate(nodes["build_presentation"].output)
    out(f"{built.presentation_id}: {len(built.slides)} slides, {built.element_count()} elements, "
        f"{built.config.aspect_ratio} {built.config.width:g}x{built.config.height:g}pt, theme={built.config.theme.name}")
    for r in built.references:
        out(f"  [{r.number}] {r.citation_id} -> {r.evidence_id} -> {r.source_id}: {_clip(r.text, 60)}")

    _heading(out, "7. Render PPTX")
    for e in container.task_service.events(task.task_id):
        if e.type.startswith(("slide_planning.", "slide_plan.", "presentation.")):
            keys = ("deck_id", "slides", "presentation_id", "renderer", "checksum", "reused", "elements")
            out(f"  {e.type:<32} " + " ".join(f"{k}={e.data[k]}" for k in keys if k in e.data))

    _heading(out, "8. Stored PRESENTATION artifact")
    art = arts["presentation"]
    meta = PresentationArtifactMetadata.model_validate(art.metadata)
    data = container.artifacts.read(art.artifact_id)
    out(f"artifact_id: {art.artifact_id}  ({art.type.value} v{art.version}, task {art.task_id})")
    out(f"media type:  {art.media_type}")
    out(f"checksum:    sha256={art.content_hash} (bytes match: {hashlib.sha256(data).hexdigest() == art.content_hash})")
    out(f"object:      {meta.object_key}  ({art.size_bytes} bytes)")
    out(f"renderer:    {meta.renderer}; slides={meta.slides} elements={meta.elements} layouts={meta.layouts}")
    out(f"parents:     {', '.join(name_of(p) for p in art.parent_ids)}")
    out(f"lineage:     {', '.join(sorted({a.name for a in container.artifacts.lineage(art.artifact_id)}))}")

    _heading(out, "9. The .pptx opened with python-pptx")
    doc = PptxDocument(io.BytesIO(data))
    checksums = {ImageAsset.model_validate(a.metadata).object.checksum: a.name for a in images}
    width_in, height_in = doc.slide_width / 914400, doc.slide_height / 914400
    out(f"{len(doc.slides)} slides, {width_in:.2f} x {height_in:.2f} in")
    for i, slide in enumerate(doc.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else "-"
        placed = [checksums.get(hashlib.sha256(s.image.blob).hexdigest(), "UNKNOWN")
                  for s in slide.shapes if s.shape_type == PICTURE]
        footer = next((s.text_frame.text for s in slide.shapes if s.name == "footer"), "")
        out(f"  {i:>2}. {_clip(title, 40)!r} pictures: {', '.join(placed) or '-'}")
        if footer:
            out(f"      footer: {_clip(footer, 90)}")
    cost = task.cost
    planner = cost.by_agent["slide_planner"]
    render = cost.by_service[f"presentation_render:{meta.renderer}"]
    out(f"\nslide planner: calls={planner.calls} tokens={planner.usage.total_tokens} ${planner.cost_usd:.6f}; "
        f"rendering: calls={render.calls} cost={'not reported' if render.cost_usd is None else render.cost_usd} "
        f"units={render.units}")
    if save_to is not None:
        save_to.write_bytes(data)
        out(f"\nsaved {save_to}")


async def run_presentation_demo(container: Container, *, out=print, save_to: Path | None = None) -> Task:
    _heading(out, "Lesson request -> ... -> Visual -> Slide Planning -> Presentation Build -> Render")
    task = await run_demo(container, out=lambda line: None)
    out(f"task {task.task_id}: {task.status.value}")
    if task.status != TaskStatus.COMPLETED:
        for err in task.errors:
            out(f"error: [{err.kind}] {err.node_id}: {err.message}")
        return task
    report_presentation(container, task, out, save_to)
    return task


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data-dir", type=Path, default=None,
                        help="where to keep the database and artifacts (default: a fresh temporary directory)")
    parser.add_argument("--out", type=Path, default=None, help="also save the generated .pptx here")
    parser.add_argument("--verbose", action="store_true", help="also print structured event logs to stderr")
    args = parser.parse_args()
    configure_logging("INFO" if args.verbose else "WARNING")

    with tempfile.TemporaryDirectory(prefix="teaching-agent-presentation-demo-") as tmp:
        data_dir = args.data_dir or Path(tmp)
        container = build_container(Settings(data_dir=data_dir))
        try:
            task = asyncio.run(run_presentation_demo(container, save_to=args.out))
        finally:
            container.close()
        if args.data_dir is None and args.out is None:
            print("\n(artifacts were written to a temporary directory; pass --out lesson.pptx to keep the file)")
    return 0 if task.status == TaskStatus.COMPLETED else 1


if __name__ == "__main__":
    raise SystemExit(main())
