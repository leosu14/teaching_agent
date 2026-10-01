"""Local PPTX renderer (python-pptx). Maps the semantic layouts to slide geometry from PresentationConfig and
PresentationTheme, places the stored IMAGE_ASSET files, and writes byte-reproducible output: the same
presentation always produces the same file, so the same checksum."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from datetime import datetime

from pptx import Presentation as PptxDocument
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

from app.providers.presentation.base import (
    PPTX_MEDIA_TYPE,
    MediaReader,
    PresentationRenderer,
    PresentationRenderError,
    RenderedPresentation,
    read_image,
)
from app.schemas.presentation import (
    ImageElement,
    Presentation,
    PresentationConfig,
    PresentationSlide,
    SlideLayout,
)

EMU_PER_PT = 12700
# Fixed timestamps (the zip format's epoch) make the output byte-reproducible, like SOURCE_DATE_EPOCH builds.
REPRODUCIBLE_TIME = (1980, 1, 1, 0, 0, 0)
TITLE_SLIDE_LAYOUT = 0  # python-pptx default template: "Title Slide"
TITLE_ONLY_LAYOUT = 5  # "Title Only"
SLIDE_NUMBER_WIDTH = 60.0
MIN_FONT = 10.0
LINE_HEIGHT = 1.2
CHAR_WIDTH = 0.5  # average glyph width as a fraction of the font size


@dataclass(frozen=True)
class Box:
    x: float
    y: float
    w: float
    h: float

    def emu(self) -> tuple[Emu, Emu, Emu, Emu]:
        return tuple(Emu(round(v * EMU_PER_PT)) for v in (self.x, self.y, self.w, self.h))  # type: ignore[return-value]

    def split_bottom(self, height: float, gap: float) -> tuple[Box, Box]:
        height = min(height, self.h / 2)
        return (Box(self.x, self.y, self.w, self.h - height - gap),
                Box(self.x, self.y + self.h - height, self.w, height))


def layout_regions(layout: SlideLayout, config: PresentationConfig) -> dict[str, Box]:
    """Semantic layout -> region boxes in points. Deterministic and driven only by the config and theme."""
    s = config.theme.spacing
    width, height = config.width, config.height
    content_w = width - 2 * s.margin
    footer_y = height - s.margin / 2 - s.footer_height
    footer = Box(s.margin, footer_y, content_w - SLIDE_NUMBER_WIDTH - s.gutter, s.footer_height)
    number = Box(width - s.margin - SLIDE_NUMBER_WIDTH, footer_y, SLIDE_NUMBER_WIDTH, s.footer_height)
    if layout == SlideLayout.TITLE:
        top = height * 0.28
        title_h = s.title_height * 1.5
        sub = Box(s.margin, top + title_h + s.block_gap, content_w, s.title_height * 1.2)
        body_y = sub.y + sub.h + s.block_gap
        return {"title": Box(s.margin, top, content_w, title_h), "subtitle": sub,
                "body": Box(s.margin, body_y, content_w, max(footer_y - s.block_gap - body_y, s.title_height)),
                "footer": footer, "number": number}

    title = Box(s.margin, s.margin, content_w, s.title_height)
    top = title.y + title.h + s.block_gap
    content = Box(s.margin, top, content_w, footer_y - s.block_gap - top)
    regions = {"title": title, "subtitle": Box(s.margin, top, content_w, s.title_height * 0.6), "body": content,
               "footer": footer, "number": number}
    half = (content.w - s.gutter) / 2
    if layout == SlideLayout.TWO_COLUMN:
        regions["left"] = Box(content.x, content.y, half, content.h)
        regions["right"] = Box(content.x + half + s.gutter, content.y, half, content.h)
    elif layout == SlideLayout.IMAGE_TEXT:
        text_w = content.w * 0.42
        regions["body"] = Box(content.x, content.y, text_w, content.h)
        regions["image"] = Box(content.x + text_w + s.gutter, content.y, content.w - text_w - s.gutter, content.h)
    elif layout == SlideLayout.FULL_IMAGE:
        regions["image"] = content
    return regions


def fit_font(texts: list[str], box: Box, size: float) -> float:
    """Shrink the font until the estimated wrapped lines fit the box (no overflow), down to MIN_FONT."""
    while size > MIN_FONT:
        chars_per_line = max(box.w / (size * CHAR_WIDTH), 1)
        lines = sum(max(1, -(-len(t) // int(chars_per_line))) for t in texts)
        if lines * size * LINE_HEIGHT <= box.h:
            break
        size -= 1
    return size


def fit_image(element: ImageElement, box: Box) -> Box:
    """Largest box with the image's aspect ratio inside `box`, centred."""
    scale = min(box.w / element.width, box.h / element.height)
    w, h = element.width * scale, element.height * scale
    return Box(box.x + (box.w - w) / 2, box.y + (box.h - h) / 2, w, h)


def _weight(element) -> float:
    match element.kind:
        case "bullets":
            return len(element.items) + 0.5
        case "table":
            return len(element.rows) + 1
        case "text":
            return max(1.0, len(element.text) / 80)
    return 1.0


def stack(box: Box, elements: list, gap: float) -> list[Box]:
    """Split a region vertically among its elements, proportionally to their estimated height."""
    weights = [_weight(e) for e in elements]
    free = box.h - gap * (len(elements) - 1)
    out, y = [], box.y
    for w in weights:
        h = free * w / sum(weights)
        out.append(Box(box.x, y, box.w, h))
        y += h + gap
    return out


def reproducible(data: bytes) -> bytes:
    """Rewrite the package with fixed member timestamps: identical presentations give identical bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            member = zipfile.ZipInfo(info.filename, date_time=REPRODUCIBLE_TIME)
            member.compress_type = zipfile.ZIP_DEFLATED
            member.external_attr = 0o644 << 16
            dst.writestr(member, src.read(info.filename))
    return out.getvalue()


class PptxPresentationRenderer(PresentationRenderer):
    name = "python-pptx"
    format = "pptx"
    media_type = PPTX_MEDIA_TYPE

    def __init__(self, media: MediaReader) -> None:
        self._media = media

    async def render(self, presentation: Presentation) -> RenderedPresentation:
        images = {e.element_id: read_image(self._media, e) for s in presentation.slides for e in s.element("image")}
        try:
            content = self._render(presentation, images)
        except PresentationRenderError:
            raise
        except Exception as exc:  # python-pptx or zip errors: the renderer failed, not the caller
            raise PresentationRenderError(f"PPTX rendering failed: {type(exc).__name__}: {exc}") from exc
        return RenderedPresentation(
            content=content, media_type=self.media_type, format=self.format, renderer=self.name,
            slides=len(presentation.slides), elements=presentation.element_count(),
            image_artifact_ids=presentation.image_artifact_ids(),
        )

    def _render(self, presentation: Presentation, images: dict[str, bytes]) -> bytes:
        config = presentation.config
        doc = PptxDocument()
        doc.slide_width = Emu(round(config.width * EMU_PER_PT))
        doc.slide_height = Emu(round(config.height * EMU_PER_PT))
        props = doc.core_properties
        props.title, props.language = presentation.title, presentation.language
        props.subject = str(presentation.metadata.get("topic", ""))
        props.author = props.last_modified_by = "teaching-agent"
        props.created = props.modified = datetime(*REPRODUCIBLE_TIME)
        props.revision = 1
        for slide in presentation.slides:
            self._slide(doc, slide, len(presentation.slides), config, images)
        buf = io.BytesIO()
        doc.save(buf)
        return reproducible(buf.getvalue())

    def _slide(self, doc, slide: PresentationSlide, total: int, config: PresentationConfig,
               images: dict[str, bytes]) -> None:
        theme = config.theme
        is_title = slide.layout == SlideLayout.TITLE
        page = doc.slides.add_slide(doc.slide_layouts[TITLE_SLIDE_LAYOUT if is_title else TITLE_ONLY_LAYOUT])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string(theme.colors.background)
        regions = layout_regions(slide.layout, config)
        by_region: dict[str, list] = {}
        for element in slide.elements:
            by_region.setdefault(element.region, []).append(element)

        captions = by_region.pop("caption", [])
        if captions:
            image_region = "right" if slide.layout == SlideLayout.TWO_COLUMN else "image"
            base = regions.get(image_region, regions["body"])
            height = theme.typography.caption * LINE_HEIGHT * 2 + theme.spacing.block_gap
            regions[image_region], regions["caption"] = base.split_bottom(height, theme.spacing.block_gap)
            by_region["caption"] = captions

        for region, elements in by_region.items():
            box = regions.get(region) or regions["body"]
            if region in ("title", "subtitle", "footer"):
                boxes = [box] * len(elements)
            else:
                boxes = stack(box, elements, theme.spacing.block_gap)
            for element, ebox in zip(elements, boxes):
                self._element(page, element, ebox, slide, config, images)

        if config.footer.show_slide_numbers:
            self._text(page, "slide_number", [f"{slide.order}/{total}"], regions["number"], config,
                       size=theme.typography.footer, color=theme.colors.muted, align=PP_ALIGN.RIGHT)
        if slide.notes:
            page.notes_slide.notes_text_frame.text = slide.notes
        if is_title and not any(e.kind == "text" and e.role == "subtitle" for e in slide.elements):
            for ph in list(page.placeholders):  # drop the template's empty subtitle placeholder
                if ph.placeholder_format.idx == 1:
                    ph._element.getparent().remove(ph._element)

    def _element(self, page, element, box: Box, slide: PresentationSlide, config: PresentationConfig,
                 images: dict[str, bytes]) -> None:
        theme = config.theme
        typo, colors = theme.typography, theme.colors
        title_layout = slide.layout == SlideLayout.TITLE
        match element.kind:
            case "title":
                shape = page.shapes.title
                self._place(shape, box)
                shape.name = element.element_id
                size = typo.title if title_layout else typo.heading
                self._fill(shape.text_frame, [element.text], fit_font([element.text], box, size), theme.fonts.heading,
                           colors.accent if not title_layout else colors.text, bold=True,
                           align=PP_ALIGN.CENTER if title_layout else PP_ALIGN.LEFT)
            case "text" if element.role == "subtitle" and title_layout:
                shape = page.placeholders[1]
                self._place(shape, box)
                shape.name = element.element_id
                self._fill(shape.text_frame, [element.text], fit_font([element.text], box, typo.subtitle),
                           theme.fonts.body, colors.muted, align=PP_ALIGN.CENTER)
            case "text":
                size = {"subtitle": typo.subtitle, "caption": typo.caption}.get(element.role, typo.body)
                color = colors.muted if element.role in ("subtitle", "caption") else (
                    colors.accent if element.role == "answer" else colors.text)
                self._text(page, element.element_id, [element.text], box, config, size=size, color=color,
                           bold=element.role == "question", align=PP_ALIGN.CENTER if title_layout else PP_ALIGN.LEFT)
            case "bullets":
                marks = [f"{i}." if element.numbered else "•" for i in range(1, len(element.items) + 1)]
                lines = [f"{m} {item}" for m, item in zip(marks, element.items)]
                self._text(page, element.element_id, lines, box, config, size=typo.body, color=colors.text)
            case "table":
                self._table(page, element, box, config)
            case "image":
                left, top, width, height = fit_image(element, box).emu()
                picture = page.shapes.add_picture(io.BytesIO(images[element.element_id]), left, top, width, height)
                picture.name = element.element_id
                picture._element.nvPicPr.cNvPr.set("descr", element.alt_text)  # alt text for screen readers
                if theme.border_width:
                    picture.line.width = Pt(theme.border_width)
                    picture.line.color.rgb = RGBColor.from_string(colors.muted)
            case "footer":
                self._text(page, "footer", [element.text], box, config, size=typo.footer, color=colors.muted)
            case _:
                raise PresentationRenderError(f"unsupported element {element.kind}")

    @staticmethod
    def _place(shape, box: Box) -> None:
        shape.left, shape.top, shape.width, shape.height = box.emu()

    def _text(self, page, name: str, lines: list[str], box: Box, config: PresentationConfig, *, size: float,
              color: str, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
        shape = page.shapes.add_textbox(*box.emu())
        shape.name = name
        self._fill(shape.text_frame, lines, fit_font(lines, box, size), config.theme.fonts.body, color, bold=bold,
                   align=align, space_after=config.theme.spacing.paragraph_after)

    @staticmethod
    def _fill(frame, lines: list[str], size: float, font: str, color: str, *, bold: bool = False,
              align=PP_ALIGN.LEFT, space_after: float = 0.0) -> None:
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.clear()
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.alignment = align
            if space_after:
                paragraph.space_after = Pt(space_after)
            run = paragraph.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.name = font
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)

    def _table(self, page, element, box: Box, config: PresentationConfig) -> None:
        theme = config.theme
        rows, cols = len(element.rows) + 1, len(element.headers)
        frame = page.shapes.add_table(rows, cols, *box.emu())
        frame.name = element.element_id
        cells = [element.headers, *element.rows]
        size = fit_font([" ".join(r) for r in cells], box, theme.typography.body * 0.8)
        for r, values in enumerate(cells):
            for c, value in enumerate(values):
                cell = frame.table.cell(r, c)
                self._fill(cell.text_frame, [value], size, theme.fonts.body,
                           "FFFFFF" if r == 0 else theme.colors.text, bold=r == 0)
